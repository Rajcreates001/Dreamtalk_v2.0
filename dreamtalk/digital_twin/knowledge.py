"""Knowledge Engine — Role-based behavior.

Knowledge behaves differently depending on user category:

NORMAL USER:
- No Knowledge Upload section in UI
- Knowledge comes automatically from: conversations, long-term memory,
  relationship, corrections, personal notes, conversation history, preferences
- Managed automatically by the Learning Engine

HEALTHCARE:
- Two independent knowledge sources:
  Source 1: Uploaded Knowledge (SOP, guidelines, papers, books, protocols)
  Source 2: Conversation Learning (transcripts, summaries, medical notes)
- Uploaded knowledge remains authoritative
- Conversation learning never overwrites validated medical knowledge without review

BUSINESS:
- Same architecture as Healthcare
- Uploaded: policies, manuals, docs, wiki, GitHub, Notion, CRM, ERP
- Conversation: meetings, customer calls, interviews, support tickets
- Company knowledge remains authoritative
"""

import json
import logging
import uuid
import os
from typing import List, Optional
from datetime import datetime
from fastapi import UploadFile

from dreamtalk.backend.db.database import execute, fetchrow, fetch
from dreamtalk.media.repository import MediaRepository

logger = logging.getLogger("dreamtalk.digital_twin.knowledge")


ALLOWED_DOC_TYPES = {
    "application/pdf", "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "text/plain", "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "text/csv",
    "application/msword", "application/vnd.ms-powerpoint", "application/vnd.ms-excel",
}


class KnowledgeIngestionPipeline:
    """Role-aware knowledge ingestion.

    - personal/normal_user: knowledge is auto-derived from conversation (no manual upload)
    - healthcare/business: accepts manual uploads + conversation learning
    """

    def __init__(self):
        self.media_repo = MediaRepository()

    async def upload_knowledge(
        self,
        twin_id: str,
        role: str,
        files: List[UploadFile],
        source_type: str = "pdf",
    ) -> dict:
        """Upload knowledge documents (healthcare/business only).

        For personal/normal_user: returns error — knowledge is auto-derived.
        """
        if role in ("personal", "normal_user"):
            return {
                "status": "rejected",
                "reason": "Personal Digital Twins do not support manual knowledge upload. "
                          "Knowledge is automatically derived from conversations.",
            }

        # Create knowledge subdirectories
        media_role = self._map_role(role)
        results = []

        for file in files:
            content = await file.read()
            if not content:
                continue

            # Store file in media repository
            asset = await self.media_repo.store_file(
                twin_id=twin_id,
                role=media_role,
                category="knowledge",
                filename=file.filename,
                content=content,
                metadata={"mime_type": file.content_type or "application/octet-stream", "source_type": source_type},
            )

            # Run ingestion pipeline
            ingestion_result = await self._ingest_document(twin_id, asset, source_type)
            results.append(ingestion_result)

        return {
            "status": "completed",
            "role": role,
            "files_processed": len(results),
            "results": results,
        }

    async def _ingest_document(self, twin_id: str, asset, source_type: str) -> dict:
        """12-step document ingestion pipeline:
        Upload → Virus Scan → OCR → Extract → Normalize → Chunk → Embed
        → Index → Knowledge Graph → Cross-Ref → Validate → Attach
        """
        result = {
            "asset_id": asset.asset_id,
            "filename": asset.filename,
            "status": "processing",
            "pipeline_steps": [],
        }

        # Step 1: Virus Scan — real ClamAV scan (or graceful pass if not installed)
        virus_ok = True
        try:
            import subprocess as _sp
            asset_file_path = getattr(asset, 'file_path', None)
            if asset_file_path and os.path.exists(asset_file_path):
                r = _sp.run(["clamscan", "--no-summary", asset_file_path], capture_output=True, text=True, timeout=30)
                if r.returncode != 0 and "OK" not in r.stdout:
                    virus_ok = False
                    result["pipeline_steps"].append({"step": "virus_scan", "status": "quarantined", "detail": r.stdout.strip()[:200]})
                else:
                    result["pipeline_steps"].append({"step": "virus_scan", "status": "passed"})
            else:
                result["pipeline_steps"].append({"step": "virus_scan", "status": "skipped", "reason": "no file path"})
        except Exception:
            result["pipeline_steps"].append({"step": "virus_scan", "status": "skipped", "reason": "ClamAV not available"})

        if not virus_ok:
            result["status"] = "quarantined"
            return result

        # Step 2: OCR (Tesseract) — for images/scanned PDFs
        extracted_text = ""
        ocr_applied = False
        try:
            import pytesseract
            from PIL import Image
            mime = getattr(asset.metadata, 'mime_type', '')
            if mime.startswith("image/") or asset.filename.lower().endswith((".png", ".jpg", ".jpeg", ".tiff", ".bmp")):
                file_path = getattr(asset, 'file_path', None)
                if file_path and os.path.exists(file_path):
                    img = Image.open(file_path)
                    extracted_text = pytesseract.image_to_string(img)
                    ocr_applied = True
                    result["pipeline_steps"].append({"step": "ocr", "status": "completed", "length": len(extracted_text)})
                else:
                    result["pipeline_steps"].append({"step": "ocr", "status": "skipped", "reason": "no file"})
            else:
                result["pipeline_steps"].append({"step": "ocr", "status": "skipped", "reason": "not an image"})
        except ImportError:
            result["pipeline_steps"].append({"step": "ocr", "status": "skipped", "reason": "pytesseract not installed"})
        except Exception as e:
            result["pipeline_steps"].append({"step": "ocr", "status": "failed", "error": str(e)[:100]})

        # Step 3: Text Extraction — real extraction for PDF/DOCX/TXT/CSV/XLSX
        ext = ''
        fp = asset.file_path if hasattr(asset, 'file_path') else ''
        try:
            if not extracted_text and fp and os.path.exists(fp):
                ext = asset.filename.lower().split('.')[-1] if '.' in asset.filename else ''
                
                if ext == 'txt':
                    with open(fp, 'r', encoding='utf-8', errors='replace') as f:
                        extracted_text = f.read()
                elif ext == 'pdf':
                    try:
                        import PyPDF2
                        with open(fp, 'rb') as f:
                            reader = PyPDF2.PdfReader(f)
                            extracted_text = '\n'.join(page.extract_text() for page in reader.pages)
                    except ImportError:
                        try:
                            from pdfminer.high_level import extract_text as pdf_extract
                            extracted_text = pdf_extract(fp)
                        except ImportError:
                            extracted_text = f"[PDF extraction requires PyPDF2 or pdfminer: {asset.filename}]"
                elif ext == 'docx':
                    try:
                        import docx
                        doc = docx.Document(fp)
                        extracted_text = '\n'.join(p.text for p in doc.paragraphs)
                    except ImportError:
                        extracted_text = f"[DOCX extraction requires python-docx: {asset.filename}]"
                elif ext == 'csv':
                    with open(fp, 'r', encoding='utf-8', errors='replace') as f:
                        extracted_text = f.read()
                elif ext in ('xlsx', 'xls'):
                    try:
                        import openpyxl
                        wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
                        rows = []
                        for sheet in wb.worksheets:
                            for row in sheet.iter_rows(values_only=True):
                                rows.append(' | '.join(str(c) if c is not None else '' for c in row))
                        extracted_text = '\n'.join(rows)
                    except ImportError:
                        extracted_text = f"[XLSX extraction requires openpyxl: {asset.filename}]"
                elif ext == 'pptx':
                    try:
                        from pptx import Presentation
                        prs = Presentation(fp)
                        slides = []
                        for slide in prs.slides:
                            for shape in slide.shapes:
                                if hasattr(shape, 'text'):
                                    slides.append(shape.text)
                        extracted_text = '\n'.join(slides)
                    except ImportError:
                        extracted_text = f"[PPTX extraction requires python-pptx: {asset.filename}]"

            result["pipeline_steps"].append({"step": "text_extraction", "status": "completed" if extracted_text else "empty", "length": len(extracted_text), "format": ext if ext else 'unknown'})
        except Exception as e:
            logger.warning("Text extraction failed for %s: %s", asset.filename, e)
            result["pipeline_steps"].append({"step": "text_extraction", "status": "failed", "error": str(e)[:100]})

        # Step 4: Normalize — clean whitespace, normalize unicode
        normalized = ""
        try:
            import unicodedata
            if extracted_text:
                normalized = unicodedata.normalize('NFKC', extracted_text)
                normalized = '\n'.join(line.strip() for line in normalized.split('\n'))
                normalized = ' '.join(normalized.split())
            result["pipeline_steps"].append({"step": "normalize", "status": "completed" if normalized else "empty", "length": len(normalized)})
        except Exception as e:
            result["pipeline_steps"].append({"step": "normalize", "status": "failed", "error": str(e)[:100]})

        # Step 5: Chunk (1000-word windows with 100-word overlap)
        chunks = self._chunk_text(normalized) if normalized else []
        result["pipeline_steps"].append({"step": "chunking", "status": "completed" if chunks else "empty", "chunks": len(chunks)})

        # Step 6-7: Embed + Index (BGE-M3 powered)
        embedding_status = "skipped"
        indexed_count = 0
        if chunks:
            try:
                from dreamtalk.digital_twin.embeddings import index_chunks
                source_id = getattr(asset, 'asset_id', None) or getattr(asset, 'id', None) or str(uuid.uuid4())
                indexed_count = await index_chunks(twin_id, source_id, chunks)
                embedding_status = "completed" if indexed_count > 0 else "failed"
            except Exception as e:
                logger.warning("Embedding failed: %s", e)
                embedding_status = f"failed: {e}"

        result["pipeline_steps"].append({"step": "embedding", "status": embedding_status, "chunks_indexed": indexed_count})

        # Step 8: Knowledge Graph — extract named entities using spaCy or simple keyword matching
        graph_entities = []
        graph_relations = []
        try:
            import spacy
            try:
                nlp = spacy.load("en_core_web_sm")
            except OSError:
                import subprocess
                subprocess.run(["python", "-m", "spacy", "download", "en_core_web_sm"], capture_output=True)
                nlp = spacy.load("en_core_web_sm")
            doc = nlp(normalized[:50000])  # Limit to 50k chars for performance
            graph_entities = list(set(ent.text for ent in doc.ents))
            graph_relations = [(tok.text, tok.dep_, tok.head.text) for tok in doc if tok.dep_ != 'punct'][:200]
            result["pipeline_steps"].append({"step": "knowledge_graph", "status": "completed", "entities": len(graph_entities), "relations": len(graph_relations)})
        except ImportError:
            # Fallback: simple keyword extraction
            import re
            words = re.findall(r'\b[A-Z][a-z]+(?: [A-Z][a-z]+)*\b', normalized[:50000])
            graph_entities = list(set(w for w in words if len(w) > 3))[:100]
            result["pipeline_steps"].append({"step": "knowledge_graph", "status": "basic", "entities": len(graph_entities)})
        except Exception as e:
            result["pipeline_steps"].append({"step": "knowledge_graph", "status": "failed", "error": str(e)[:100]})

        # Step 9: Cross-Reference — link to existing knowledge in the same twin
        cross_refs = []
        try:
            existing = await fetch(
                "SELECT id, content FROM knowledge_chunks WHERE twin_id = $1 ORDER BY chunk_index LIMIT 50",
                twin_id,
            )
            if existing and chunks:
                import difflib
                for i, chunk in enumerate(chunks[:10]):
                    for row in existing:
                        ratio = difflib.SequenceMatcher(None, chunk[:200], row['content'][:200]).ratio()
                        if ratio > 0.6:
                            cross_refs.append({"source_chunk": i, "target_id": row['id'], "similarity": round(ratio, 3)})
            result["pipeline_steps"].append({"step": "cross_reference", "status": "completed" if cross_refs else "none_found", "references": len(cross_refs)})
        except Exception as e:
            result["pipeline_steps"].append({"step": "cross_reference", "status": "failed", "error": str(e)[:100]})

        # Step 10: Validate — quality check on extracted content
        validation = {"has_content": bool(extracted_text), "min_length": len(extracted_text) >= 50, "has_meaningful_chunks": len(chunks) > 0}
        validation_score = sum(1 for v in validation.values() if v) / len(validation)
        try:
            result["pipeline_steps"].append({"step": "validation", "status": "passed" if validation_score >= 0.66 else "low_quality", "score": validation_score, "details": validation})
        except Exception as e:
            result["pipeline_steps"].append({"step": "validation", "status": "failed", "error": str(e)[:100]})

        # Step 11: Store in knowledge_sources + knowledge_chunks tables
        try:
            await self._store_knowledge(twin_id, asset, source_type, chunks)
            result["pipeline_steps"].append({"step": "attach", "status": "completed", "chunks_stored": len(chunks)})
        except Exception as e:
            result["pipeline_steps"].append({"step": "attach", "status": "failed", "error": str(e)[:100]})

        result["status"] = "completed"
        result["chunk_count"] = len(chunks)
        result["entities_found"] = len(graph_entities)
        result["cross_references"] = len(cross_refs)

        return result

    def _chunk_text(self, text: str, chunk_size: int = 1000, overlap: int = 100) -> List[str]:
        if not text:
            return []
        words = text.split()
        chunks = []
        start = 0
        while start < len(words):
            end = min(start + chunk_size, len(words))
            chunk = " ".join(words[start:end])
            if chunk.strip():
                chunks.append(chunk)
            start += chunk_size - overlap
        return chunks

    async def _store_knowledge(self, twin_id: str, asset, source_type: str, chunks: List[str]):
        """Store knowledge source and chunks in database."""
        source_id = str(uuid.uuid4())

        # Insert knowledge source
        await execute(
            """INSERT INTO knowledge_sources (id, twin_id, source_type, source_name, original_filename,
               file_path, file_size, mime_type, status, chunk_count)
               VALUES ($1, $2, $3, $4, $5, $6, $7, $8, $9, $10)""",
            source_id, twin_id, source_type, asset.filename, asset.filename,
            asset.file_path, asset.metadata.file_size, asset.metadata.mime_type,
            "ingested", len(chunks),
        )

        # Insert chunks
        for i, chunk_text in enumerate(chunks):
            await execute(
                """INSERT INTO knowledge_chunks (source_id, twin_id, chunk_index, content)
                   VALUES ($1, $2, $3, $4)""",
                source_id, twin_id, i, chunk_text,
            )

    async def get_knowledge_sources(self, twin_id: str) -> List[dict]:
        rows = await fetch(
            "SELECT id, source_type, source_name, original_filename, status, chunk_count, created_at "
            "FROM knowledge_sources WHERE twin_id = $1 ORDER BY created_at DESC",
            twin_id,
        )
        return [dict(r) for r in rows]

    async def delete_knowledge_source(self, source_id: str) -> bool:
        await execute("DELETE FROM knowledge_sources WHERE id = $1", source_id)
        return True

    def _map_role(self, role: str) -> str:
        if role in ("personal", "normal_user"):
            return "normal_user"
        elif role == "healthcare":
            return "healthcare"
        elif role == "business":
            return "business"
        return "normal_user"
